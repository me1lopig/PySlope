import io
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_report(layers_df, water_df, loads_df, seismic_params, calc_params, results, fig_bytes):
    doc = Document()
    
    title = doc.add_heading('Memoria de Cálculo: Estabilidad de Taludes', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Análisis de equilibrio límite 2D mediante método de Bishop Simplificado.').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading('1. Resultados del Análisis', level=1)
    if results['fs_min'] != float('inf'):
        p = doc.add_paragraph()
        p.add_run(f"Factor de Seguridad Mínimo (FS): {results['fs_min']:.3f}\n").bold = True
        p.add_run(f"Centro Crítico (Xc, Yc): ({results['xc']:.2f}, {results['yc']:.2f}) m\n")
        p.add_run(f"Radio de Rotura Crítico (R): {results['r']:.2f} m")
    else:
        doc.add_paragraph("No se encontró superficie de rotura cinemáticamente admisible en la malla explorada.")
    
    doc.add_heading('2. Sección Analizada y Superficie Crítica', level=1)
    image_stream = io.BytesIO(fig_bytes)
    doc.add_picture(image_stream, width=Inches(6.0))
    doc.add_paragraph('Figura 1. Geometría del modelo, estratigrafía y círculo crítico de deslizamiento.').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading('3. Modelo Geotécnico', level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    headers = ['Nombre', 'Gamma (kN/m3)', 'c (kPa)', 'phi (º)']
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        
    for _, row in layers_df.iterrows():
        row_cells = table.add_row().cells
        row_cells[0].text = str(row['Nombre'])
        row_cells[1].text = f"{row['Gamma (kN/m3)']:.1f}"
        row_cells[2].text = f"{row['c (kPa)']:.1f}"
        row_cells[3].text = f"{row['phi (º)']:.1f}"

    doc.add_heading('4. Condiciones de Contorno y Sismo', level=1)
    doc.add_paragraph(f"Coeficiente sísmico horizontal (kh): {seismic_params['kh']:.2f} g")
    doc.add_paragraph(f"Coeficiente sísmico vertical (kv): {seismic_params['kv']:.2f} g")
    
    if not loads_df.empty:
        doc.add_heading('Sobrecargas en Superficie', level=2)
        for _, row in loads_df.iterrows():
            doc.add_paragraph(f"Carga distribuida q = {row['q (kPa)']} kPa entre X={row['X inicio (m)']} m y X={row['X fin (m)']} m.")
    
    doc.add_heading('5. Configuración de la Malla de Cálculo', level=1)
    doc.add_paragraph(f"Dovelas en discretización: {calc_params['n_slices']}")
    doc.add_paragraph(f"Resolución de malla (X, Y, Radios): {calc_params['gnx']} x {calc_params['gny']} x {calc_params['gnr']}")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    return buffer
